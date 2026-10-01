"""
=========================================================
DOCX Reader

Reads Microsoft Word (.docx) documents.

Supports:
- Paragraphs
- Tables

This reader is part of the AI Engine.
=========================================================
"""

from docx import Document

from ai_engine.readers.base_reader import BaseReader


class DocxReader(BaseReader):
    """
    Reader for DOCX documents.
    """

    def read(self, file_path: str) -> str:
        """
        Read a DOCX file and extract all readable text.
        """

        document = Document(file_path)

        content = []

        # ==========================================
        # Read normal paragraphs
        # ==========================================

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                content.append(text)

        # ==========================================
        # Read tables
        # ==========================================

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    text = cell.text.strip()

                    if text:

                        cells.append(text)

                if cells:

                    content.append(" | ".join(cells))

        return "\n".join(content)